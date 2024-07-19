import sys
import os
import pandas as pd
from azure.identity import DefaultAzureCredential, InteractiveBrowserCredential
from azure.mgmt.resource import ResourceManagementClient
from openpyxl.styles import PatternFill, Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from datetime import datetime
from PyQt5.QtWidgets import QApplication, QWidget, QVBoxLayout, QPushButton, QTextEdit, QProgressBar
from PyQt5.QtCore import Qt, QThread, pyqtSignal

class ExportThread(QThread):
    update_progress = pyqtSignal(int)
    log_message = pyqtSignal(str)

    def __init__(self):
        super().__init__()
        self.subscription_ids = [
            '7f9b36c9-5325-495c-aa2a-7c7350f302a0', 'f141e7fc-929b-4ae9-a968-94222e866042',
            'ee1c5b76-97f0-46ef-bbb0-62e6f080456b', '692efea0-2038-48e2-bf74-a333fef44d58',
            '7bd7cad1-cba0-45dd-8962-6efa16660e26', 'f18b2c7d-1eb0-46d3-9ebf-1d5073e05239'
        ]
        self.subscription_map = {
            '7f9b36c9-5325-495c-aa2a-7c7350f302a0': 'bmgf-eds-nonprod-00',
            'f141e7fc-929b-4ae9-a968-94222e866042': 'bmgf-eds-nonprod-traditional-00',
            'ee1c5b76-97f0-46ef-bbb0-62e6f080456b': 'bmgf-eds-prod-00',
            '692efea0-2038-48e2-bf74-a333fef44d58': 'Global Data & Analytics - NonProd',
            '7bd7cad1-cba0-45dd-8962-6efa16660e26': 'Global Data & Analytics - Prod',
            'f18b2c7d-1eb0-46d3-9ebf-1d5073e05239': 'ISS Azure PROD'
        }
        self.location_map = {
            'eastus': 'East US', 'westus': 'West US', 'northeurope': 'North Europe', 'australiaeast': 'Australia East'
        }
        self.type_map = {
            'microsoft.insights/components': 'Application Insights',
            'microsoft.cache/redis': 'Azure Cache for Redis',
            'microsoft.containerregistry/registries': 'Container registry'
        }
        self.excluded_types = [
              'dell.storage/filesystems', 
    'microsoft.cdn/profiles/customdomains', 
    'microsoft.sovereign/landingzoneconfigurations', 
    'microsoft.hardwaresecuritymodules/cloudhsmclusters', 
    'microsoft.cloudtest/accounts', 
    'microsoft.cloudtest/hostedpools', 
    'microsoft.cloudtest/images', 
    'microsoft.cloudtest/pools', 
    'microsoft.compute/computefleetinstances', 
    'microsoft.compute/standbypoolinstance', 
    'microsoft.compute/virtualmachineflexinstances', 
    'microsoft.kubernetesconfiguration/extensions', 
    'microsoft.containerservice/managedclusters/microsoft.kubernetesconfiguration/extensions', 
    'microsoft.kubernetes/connectedclusters/microsoft.kubernetesconfiguration/namespaces', 
    'microsoft.containerservice/managedclusters/microsoft.kubernetesconfiguration/namespaces', 
    'microsoft.kubernetes/connectedclusters/microsoft.kubernetesconfiguration/fluxconfigurations', 
    'microsoft.containerservice/managedclusters/microsoft.kubernetesconfiguration/fluxconfigurations', 
    'microsoft.portalservices/extensions/deployments', 
    'microsoft.portalservices/extensions', 
    'microsoft.portalservices/extensions/slots', 
    'microsoft.portalservices/extensions/versions', 
    'microsoft.datacollaboration/workspaces', 
    'microsoft.deviceregistry/devices', 
    'microsoft.deviceupdate/updateaccounts/activedeployments', 
    'microsoft.deviceupdate/updateaccounts/agents', 
    'microsoft.deviceupdate/updateaccounts/deployments', 
    'microsoft.deviceupdate/updateaccounts/deviceclasses', 
    'microsoft.deviceupdate/updateaccounts/updates', 
    'microsoft.deviceupdate/updateaccounts', 
    'microsoft.devopsinfrastructure/pools', 
    'microsoft.network/dnsresolverdomainlists', 
    'microsoft.network/dnsresolverpolicies', 
    'microsoft.impact/connectors', 
    'microsoft.edgeorder/virtual_orderitems', 
    'microsoft.workloads/epicvirtualinstances', 
    'microsoft.fairfieldgardens/provisioningresources/provisioningpolicies', 
    'microsoft.fairfieldgardens/provisioningresources', 
    'microsoft.fileshares/fileshares', 
    'microsoft.healthmodel/healthmodels', 
    'microsoft.hybridcompute/arcserverwithwac', 
    'microsoft.hybridcompute/machinessovereign', 
    'microsoft.hybridcompute/machinesesu', 
    'microsoft.network/virtualhubs', 
    'microsoft.network/networkvirtualappliances', 
    'microsoft.modsimworkbench/workbenches/chambers', 
    'microsoft.modsimworkbench/workbenches/chambers/connectors', 
    'microsoft.modsimworkbench/workbenches/chambers/files', 
    'microsoft.modsimworkbench/workbenches/chambers/filerequests', 
    'microsoft.modsimworkbench/workbenches/chambers/licenses', 
    'microsoft.modsimworkbench/workbenches/chambers/storages', 
    'microsoft.modsimworkbench/workbenches/chambers/workloads', 
    'microsoft.modsimworkbench/workbenches/sharedstorages', 
    'microsoft.insights/diagnosticsettings', 
    'microsoft.network/serviceendpointpolicies', 
    'microsoft.resources/resourcegraphvisualizer', 
    'microsoft.openlogisticsplatform/workspaces', 
    'microsoft.iotoperationsmq/mq', 
    'microsoft.orbital/cloudaccessrouters', 
    'microsoft.orbital/terminals', 
    'microsoft.orbital/sdwancontrollers', 
    'microsoft.recommendationsservice/accounts/modeling', 
    'microsoft.recommendationsservice/accounts/serviceendpoints', 
    'microsoft.recoveryservicesbvtd/vaults', 
    'microsoft.recoveryservicesbvtd2/vaults', 
    'microsoft.recoveryservicesintd/vaults', 
    'microsoft.recoveryservicesintd2/vaults', 
    'microsoft.features/featureprovidernamespaces/featureconfigurations', 
    'microsoft.deploymentmanager/rollouts', 
    'microsoft.providerhub/providerregistrations', 
    'microsoft.providerhub/providerregistrations/customrollouts', 
    'microsoft.providerhub/providerregistrations/defaultrollouts', 
    'microsoft.datareplication/replicationvaults', 
    'microsoft.synapse/workspaces/sqlpools', 
    'microsoft.mission/catalogs', 
    'microsoft.mission/communities', 
    'microsoft.mission/communities/communityendpoints', 
    'microsoft.mission/enclaveconnections', 
    'microsoft.mission/virtualenclaves/enclaveendpoints', 
    'microsoft.mission/virtualenclaves/endpoints', 
    'microsoft.mission/externalconnections', 
    'microsoft.mission/internalconnections', 
    'microsoft.mission/communities/transithubs', 
    'microsoft.mission/virtualenclaves', 
    'microsoft.mission/virtualenclaves/workloads', 
    'microsoft.windowspushnotificationservices/registrations', 
    'microsoft.workloads/insights', 
    'microsoft.hanaonazure/sapmonitors', 
    'microsoft.cloudhealth/healthmodels', 
    'microsoft.connectedcache/enterprisemcccustomers/enterprisemcccachenodes', 
    'microsoft.manufacturingplatform/manufacturingdataservices', 
    'microsoft.windowsesu/multipleactivationkeys', 
    'microsoft.sql/servers/databases', 
    'microsoft.sql/servers'
        ]

    def authenticate(self):
        try:
            self.log_message.emit("Attempting to authenticate using DefaultAzureCredential...")
            credential = DefaultAzureCredential()
            credential.get_token("https://management.azure.com/.default")
        except Exception as e:
            self.log_message.emit("DefaultAzureCredential failed. Falling back to InteractiveBrowserCredential.")
            self.log_message.emit(str(e))
            credential = InteractiveBrowserCredential()
        return credential

    def get_resource_group_from_id(self, resource_id):
        parts = resource_id.split('/')
        try:
            return parts[parts.index('resourceGroups') + 1]
        except ValueError:
            return None

    def run(self):
        self.log_message.emit("Initializing credentials and clients...")
        credential = self.authenticate()

        processed_data = []
        total_resources = 0

        for subscription_id in self.subscription_ids:
            resource_client = ResourceManagementClient(credential, subscription_id)
            resources = list(resource_client.resources.list())
            total_resources += len(resources)

        self.update_progress.emit(0)
        self.log_message.emit("In progress...")

        current_resource = 0
        for subscription_id in self.subscription_ids:
            resource_client = ResourceManagementClient(credential, subscription_id)
            for item in resource_client.resources.list():
                if item.type not in self.excluded_types:
                    processed_data.append({
                        "SUBSCRIPTION_NAME": self.subscription_map.get(subscription_id, subscription_id),
                        "SUBSCRIPTION_ID": subscription_id,
                        "RESOURCE_GROUP": self.get_resource_group_from_id(item.id),
                        "RESOURCE_NAME": item.name,
                        "LOCATION": self.location_map.get(item.location, item.location),
                        "TYPE": self.type_map.get(item.type, item.type),
                        "KIND": item.kind,
                        "TAG_APPLICATION": item.tags.get("application") if item.tags else None,
                        "TAG_OWNER": item.tags.get("owner") if item.tags else None,
                        "TAG_COST_CENTER": item.tags.get("cost-center") if item.tags else None,
                        "TAG_ENVIRONMENT": item.tags.get("environment") if item.tags else None,
                        "ID": item.id
                    })
                    current_resource += 1
                    self.update_progress.emit(current_resource)

        self.log_message.emit("Processing data...")

        df = pd.DataFrame(processed_data).sort_values(by=['TAG_APPLICATION', 'SUBSCRIPTION_NAME'], key=lambda x: x.str.lower())
        output_file = f"{datetime.now().strftime('%B_%d_%Y')}_EDS_DIO_AZ_Resources_Inventory.xlsx"

        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name="Pulled_Azure_Resources")
            worksheet = writer.sheets["Pulled_Azure_Resources"]
            ref = f"A1:{chr(65 + df.shape[1] - 1)}{df.shape[0] + 1}"
            table = Table(displayName="Table1", ref=ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium15", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=True)
            worksheet.add_table(table)
            header_fill = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
            header_font = Font(color="FFFFFF")
            for cell in worksheet["1:1"]:
                cell.fill = header_fill
                cell.font = header_font
            for col in worksheet.columns:
                max_length = max(len(str(cell.value)) for cell in col if cell.value is not None)
                worksheet.column_dimensions[col[0].column_letter].width = max_length + 2

        self.update_progress.emit(total_resources)
        self.log_message.emit(f"Data has been successfully exported to {output_file}")
        self.log_message.emit("Export complete!")

class AzureResourceApp(QWidget):
    def __init__(self):
        super().__init__()
        self.initUI()
        self.setup_logging()

    def initUI(self):
        self.setGeometry(100, 100, 700, 700)
        self.center()
        layout = QVBoxLayout()

        self.exportButton = QPushButton('Export Azure Resources to Excel', self)
        self.exportButton.clicked.connect(self.start_export)
        layout.addWidget(self.exportButton)

        self.progress = QProgressBar(self)
        self.progress.setAlignment(Qt.AlignCenter)
        layout.addWidget(self.progress)

        self.log = QTextEdit(self)
        self.log.setReadOnly(True)
        self.log.setFontPointSize(12)  # Set the font size here
        layout.addWidget(self.log)

        self.setLayout(layout)
        self.setWindowTitle('Azure Resource Exporter')
        self.show()

    def setup_logging(self):
        if not os.path.exists('logs'):
            os.makedirs('logs')
        self.log_file = datetime.now().strftime("logs/log_%Y_%m_%d_%H_%M_%S.txt")

    def log_message(self, message):
        timestamped_message = f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')} - {message}"
        self.log.append(timestamped_message)
        with open(self.log_file, 'a') as log_file:
            log_file.write(timestamped_message + '\n')

    def center(self):
        frame_geometry = self.frameGeometry()
        screen = QApplication.desktop().screenNumber(QApplication.desktop().cursor().pos())
        center_point = QApplication.desktop().screenGeometry(screen).center()
        frame_geometry.moveCenter(center_point)
        self.move(frame_geometry.topLeft())

    def start_export(self):
        self.thread = ExportThread()
        self.thread.update_progress.connect(self.progress.setValue)
        self.thread.log_message.connect(self.log_message)
        self.thread.start()

if __name__ == '__main__':
    app = QApplication(sys.argv)
    ex = AzureResourceApp()
    sys.exit(app.exec_())
